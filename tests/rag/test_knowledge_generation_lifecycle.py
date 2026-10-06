from __future__ import annotations

import json
import math
import os
import time
import tracemalloc
from pathlib import Path

from docx import Document as WordDocument

from backend.rag.chunking import StructureAwareChunker
from backend.rag.config import (
    RagChunkingConfig,
    RagConfig,
    RagEmbeddingConfig,
    RagRetrievalConfig,
    RagVectorStoreConfig,
)
from backend.rag.index_audit import audit_index_consistency
from backend.rag.index_manifest import IndexManifest, IndexStatus
from backend.rag.index_service import IndexService
from backend.rag.models import DocumentChunk
from backend.rag.parsers import parse_document
from backend.rag.retrieval_service import RetrievalService
from backend.rag.sparse import BM25SparseRetriever
from backend.rag.stores import FaissVectorStore, VectorSearchFilter
from backend.services.knowledge_library_service import KnowledgeLibraryService


class FixtureEmbedding:
    model_name = "lifecycle-fixture-v1"
    dimension = 4

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [[1.0, 0.0, 0.0, 0.0] for _ in texts]

    def embed_query(self, _query: str) -> list[float]:
        return [1.0, 0.0, 0.0, 0.0]


def _build_services(root: Path, allowed_root: Path):
    vector_config = RagVectorStoreConfig(
        storage_path=str(root / "faiss"),
        collection_name="knowledge_lifecycle_test",
    )
    chunking = RagChunkingConfig(
        target_tokens=64,
        overlap_tokens=4,
        minimum_tokens=3,
    )
    config = RagConfig(
        chunking=chunking,
        embedding=RagEmbeddingConfig(
            model="lifecycle-fixture-v1",
            dimension=4,
            device="cpu",
        ),
        vector_store=vector_config,
    )
    embedding = FixtureEmbedding()
    vector_store = FaissVectorStore(vector_config, dimension=4)
    sparse = BM25SparseRetriever(root / "bm25.json")
    manifest = IndexManifest(root / "manifest.json")
    manifest.recover_interrupted_operations()
    index = IndexService(
        chunker=StructureAwareChunker(chunking),
        embedding_provider=embedding,
        vector_store=vector_store,
        sparse_retriever=sparse,
        manifest=manifest,
        parser=parse_document,
    )
    library = KnowledgeLibraryService(
        index_service=index,
        manifest=manifest,
        config=config,
        embedding_provider=embedding,
        allowed_roots=(allowed_root,),
    )
    retrieval = RetrievalService(
        embedding_provider=embedding,
        vector_store=vector_store,
        sparse_retriever=sparse,
        config=RagRetrievalConfig(
            dense_top_k=10,
            sparse_top_k=10,
            fusion_top_k=10,
            final_top_k=5,
        ),
        manifest=manifest,
    )
    return library, vector_store, sparse, manifest, retrieval, config


def _write_minimal_searchable_pdf(path: Path, text: str) -> None:
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(content)).encode("ascii")
        + b" >>\nstream\n"
        + content
        + b"\nendstream",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{number} 0 obj\n".encode("ascii"))
        payload.extend(body)
        payload.extend(b"\nendobj\n")
    xref_offset = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    path.write_bytes(payload)


def _write_source_files(root: Path) -> dict[str, tuple[Path, str]]:
    root.mkdir(parents=True, exist_ok=True)
    paths: dict[str, tuple[Path, str]] = {}

    pdf_path = root / "paper.pdf"
    pdf_term = "quartzpdf7319"
    _write_minimal_searchable_pdf(pdf_path, f"{pdf_term} source evidence")
    paths["pdf"] = (pdf_path, pdf_term)

    docx_path = root / "paper.docx"
    docx_term = "quartzdocx7319"
    word = WordDocument()
    word.add_heading("DOCX lifecycle fixture", level=1)
    word.add_paragraph(f"{docx_term} source evidence")
    word.save(docx_path)
    paths["docx"] = (docx_path, docx_term)

    html_path = root / "paper.html"
    html_term = "quartzhtml7319"
    html_path.write_text(
        "<!doctype html><html><head><title>HTML lifecycle fixture</title></head>"
        f"<body><h1>Evidence</h1><p>{html_term} source evidence</p></body></html>",
        encoding="utf-8",
    )
    paths["html"] = (html_path, html_term)

    text_path = root / "paper.txt"
    text_term = "quartztxt7319"
    text_path.write_text(f"{text_term} source evidence\n", encoding="utf-8")
    paths["txt"] = (text_path, text_term)
    return paths


def _generation_chunk_ids(chunks: list[DocumentChunk]) -> set[tuple[str, str, str]]:
    return {
        (
            chunk.document_id,
            str(chunk.metadata.get("index_generation", "")),
            chunk.chunk_id,
        )
        for chunk in chunks
    }


def test_real_format_import_reindex_restart_and_delete_lifecycle(
    tmp_path: Path,
) -> None:
    allowed_root = tmp_path / "allowed"
    sources = _write_source_files(allowed_root)
    state_root = tmp_path / "state"
    (
        library,
        vector_store,
        sparse,
        manifest,
        retrieval,
        config,
    ) = _build_services(state_root, allowed_root)
    vector_store_open = True

    import_latencies_ms: list[float] = []
    retrieval_latencies_ms: list[float] = []
    import_count = 0
    tracemalloc.start()
    try:
        document_ids: dict[str, str] = {}
        active_generation_ids: dict[str, str] = {}
        for format_name, (path, _term) in sources.items():
            started = time.perf_counter()
            result = library.import_document(path)
            import_latencies_ms.append((time.perf_counter() - started) * 1000)
            import_count += 1
            assert result.status is IndexStatus.READY
            record = library.get_document(result.document_id)
            assert record is not None
            assert record.generation_id
            document_ids[format_name] = result.document_id
            active_generation_ids[format_name] = record.generation_id

        before_reindex = library.get_document(document_ids["txt"])
        assert before_reindex is not None
        old_txt_generation = before_reindex.generation_id
        text_path = sources["txt"][0]
        new_txt_term = "quartzrefresh7320"
        text_path.write_text(f"{new_txt_term} refreshed evidence\n", encoding="utf-8")
        started = time.perf_counter()
        reindexed = library.import_document(text_path)
        import_latencies_ms.append((time.perf_counter() - started) * 1000)
        import_count += 1
        assert reindexed.status is IndexStatus.READY
        new_txt_record = library.get_document(document_ids["txt"])
        assert new_txt_record is not None
        assert new_txt_record.generation_id != old_txt_generation
        old_txt_record = manifest.get_generation(
            document_ids["txt"], old_txt_generation
        )
        assert old_txt_record is not None
        assert old_txt_record.status.value == "retired"
        active_generation_ids["txt"] = new_txt_record.generation_id

        expected_stored_ids = {
            (record.document_id, generation.generation_id, chunk_id)
            for record in manifest.list_records()
            for generation in manifest.list_generations(record.document_id)
            if generation.status.value in {"ready", "retired"}
            for chunk_id in generation.chunk_ids
        }
        assert _generation_chunk_ids(vector_store.list_chunks()) == expected_stored_ids
        assert _generation_chunk_ids(sparse.list_chunks()) == expected_stored_ids
        audit = audit_index_consistency(
            manifest=manifest,
            sparse_retriever=sparse,
            vector_store=vector_store,
        )
        assert audit.status == "consistent"
        assert audit.chunk_ids_equal is True

        for format_name, (_path, term) in sources.items():
            query_term = new_txt_term if format_name == "txt" else term
            filters = VectorSearchFilter(document_ids=[document_ids[format_name]])
            for _ in range(5):
                started = time.perf_counter()
                result = retrieval.retrieve(
                    query_term,
                    filters=filters,
                    reranker_enabled=False,
                    small_to_big_enabled=False,
                )
                retrieval_latencies_ms.append((time.perf_counter() - started) * 1000)
                assert result.candidates
                assert all(
                    candidate.chunk.metadata.get("index_generation")
                    == active_generation_ids[format_name]
                    for candidate in result.candidates
                )
                assert any(
                    query_term in candidate.chunk.text
                    for candidate in result.candidates
                )

        vector_path = Path(config.vector_store.storage_path)
        vector_store.close()
        vector_store_open = False
        (
            library,
            vector_store,
            sparse,
            manifest,
            retrieval,
            _config,
        ) = _build_services(state_root, allowed_root)
        vector_store_open = True
        for format_name, (_path, term) in sources.items():
            query_term = new_txt_term if format_name == "txt" else term
            result = retrieval.retrieve(
                query_term,
                filters=VectorSearchFilter(document_ids=[document_ids[format_name]]),
                reranker_enabled=False,
                small_to_big_enabled=False,
            )
            assert result.candidates
            assert any(query_term in item.chunk.text for item in result.candidates)

        for document_id in document_ids.values():
            assert library.delete_document(document_id) is True
            assert library.delete_document(document_id) is False
        assert library.list_documents() == []
        assert vector_store.count_chunks() == 0
        assert sparse.list_chunks() == []
        assert vector_store.list_chunks() == []
        assert all(
            manifest.list_generations(document_id) == []
            for document_id in document_ids.values()
        )
        vector_store.close()
        vector_store_open = False

        # Reopen each persistent store after deletion to verify durable cleanup.
        reopened_vector = FaissVectorStore(
            RagVectorStoreConfig(
                storage_path=str(vector_path),
                collection_name="knowledge_lifecycle_test",
            ),
            dimension=4,
        )
        vector_store = reopened_vector
        vector_store_open = True
        reopened_sparse = BM25SparseRetriever(state_root / "bm25.json")
        reopened_manifest = IndexManifest(state_root / "manifest.json")
        reopened_manifest.recover_interrupted_operations()
        assert reopened_vector.count_chunks() == 0
        assert reopened_vector.list_chunks() == []
        assert reopened_sparse.list_chunks() == []
        assert reopened_manifest.list_records() == []
        reopened_vector.close()
        vector_store_open = False
    finally:
        if tracemalloc.is_tracing():
            _current_bytes, peak_bytes = tracemalloc.get_traced_memory()
            tracemalloc.stop()
        else:
            peak_bytes = 0
        if vector_store_open:
            vector_store.close()

    if os.getenv("AITRANS_RAG_LIFECYCLE_METRICS"):
        import_p95 = sorted(import_latencies_ms)[
            max(0, math.ceil(len(import_latencies_ms) * 0.95) - 1)
        ]
        retrieval_p95 = sorted(retrieval_latencies_ms)[
            max(0, math.ceil(len(retrieval_latencies_ms) * 0.95) - 1)
        ]
        print(
            json.dumps(
                {
                    "formats": sorted(sources),
                    "index_operations": import_count,
                    "index_ops_per_second": round(
                        import_count / (sum(import_latencies_ms) / 1000), 3
                    ),
                    "index_operation_p95_ms": round(import_p95, 3),
                    "retrieval_requests": len(retrieval_latencies_ms),
                    "retrieval_p95_ms": round(retrieval_p95, 3),
                    "python_tracemalloc_peak_mib": round(peak_bytes / (1024 * 1024), 3),
                    "environment_note": "synthetic four-file smoke; not production SLO",
                },
                ensure_ascii=False,
            )
        )
